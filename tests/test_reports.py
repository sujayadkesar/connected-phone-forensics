"""Recovered-data views, the Excel workbook and the PDF summary, on the synthetic demo case."""

from datetime import datetime
from pathlib import Path

import pytest

from connected_phone_forensics import content
from connected_phone_forensics.casedata import CaseData
from connected_phone_forensics.demo import build_demo_evidence
from connected_phone_forensics.engine import scan
from connected_phone_forensics.output import write_all


@pytest.fixture(scope="module")
def demo_case(tmp_path_factory):
    base = tmp_path_factory.mktemp("reports")
    result = scan(build_demo_evidence(base / "demo"), case={"case": "2026-0142", "examiner": "A. Examiner"},
                  export_dir=base / "out" / "exports")
    written = write_all(result, base / "out")
    return CaseData.load(base / "out"), written


def test_recovered_data_is_sorted_into_kinds(demo_case):
    case, _ = demo_case
    items = content.extract(case.records, case.devices)
    n = {k: len(v) for k, v in items.items()}
    assert n["messages"] == 9 and n["calls"] == 2 and n["contacts"] == 2 and n["notifications"] == 2
    assert n["media"] == 4 and n["apps"] == 6 and n["connections"] > 20
    sms = next(m for m in items["messages"] if m["text"] == "Are we still on for tonight?")
    assert sms["party"] == "Jordan Lee (+14155550177)" and sms["direction"] == "received" and sms["service"] == "SMS"
    deleted = [m for m in items["messages"] if m["status"]]
    assert {m["text"] for m in deleted} == {"Bring the documents. Delete this.", "Wipe the laptop tonight"}
    assert all(m["status"] == "Deleted, recovered" for m in deleted)
    call = next(c for c in items["calls"] if c["direction"] == "outgoing")
    assert call["party"] == "Jordan Lee (+14155550177)" and call["duration"] == 184
    photo = next(m for m in items["media"] if m["kind"] == "Photo")
    assert photo["file"] and (case.outdir / "exports" / photo["file"]).is_file()
    assert {a["package"] for a in items["apps"]} >= {"com.whatsapp", "org.whispersystems.signal"}
    for rows in items.values():  # every item points back at its record
        for r in rows:
            assert case.records[r["record"]]["source_path"] == r["source_path"]


def test_workbook_has_a_summary_and_a_sheet_per_phone(demo_case):
    from openpyxl import load_workbook

    case, written = demo_case
    path = case.outdir / "workbook.xlsx"
    assert path in written
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "Summary" and len(wb.sheetnames) == 1 + len(case.phones) + 1
    assert "Hotspots" in wb.sheetnames and case.hotspots  # the demo's iPhone hotspot, as a network
    summary = wb["Summary"]
    assert {"Summary_glance", "Summary_phones", "Summary_computers"} <= set(summary.tables)
    head = next(c for c in summary["A"] if c.value == "Recovered")  # the at-a-glance table's header
    assert head.font.b and head.alignment.horizontal == "center" and head.fill.fgColor.rgb.endswith("2456DE")
    phones = summary.tables["Summary_phones"]  # a Table object
    first_row = int(phones.ref.split(":")[0][1:]) + 1
    link = next(c for c in summary[first_row] if c.value == "Open sheet")
    assert link.hyperlink.location or link.hyperlink.target
    sheet = wb[next(n for n in wb.sheetnames if n.startswith("Phone line +14155550199"))]
    messages = sheet.tables[next(name for name in sheet.tables if name.endswith("_messages"))]
    top = int(messages.ref.split(":")[0][1:])
    assert sheet.cell(row=top, column=1).value == "Time (UTC)"
    assert isinstance(sheet.cell(row=top + 1, column=1).value, datetime)  # real date-times sort and filter


def test_workbook_never_turns_evidence_into_formulas(demo_case, tmp_path):
    from openpyxl import load_workbook

    from connected_phone_forensics.output.excel import write_workbook

    case, _ = demo_case
    case = CaseData.load(case.outdir)
    rec = next(r for r in case.records if r.get("event") == "sms_received" and r["artifact"] == "win_phonelink")
    rec["data"]["body"] = "=HYPERLINK(\"http://x\",\"click\")\x07"
    out = write_workbook(case, tmp_path / "w.xlsx")
    wb = load_workbook(out)
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert cells and all(c.data_type == "s" for c in cells)
    assert all(c.value.endswith("\\x07") for c in cells)


def _text(page) -> str:
    # Qt's headless ("offscreen") font engine writes spaces as tabs in the text layer
    return " ".join(page.extract_text().split())


def test_pdf_summary_lists_every_phone(demo_case):
    pypdf = pytest.importorskip("pypdf")
    case, written = demo_case
    path = case.outdir / "summary.pdf"
    assert path in written
    reader = pypdf.PdfReader(str(path))
    n = len(reader.pages)
    assert n >= 2 + len(case.phones)
    text = "\n".join(_text(p) for p in reader.pages)
    assert "2026-0142" in text and f"Page 1 of {n}" in text and f"Page {n} of {n}" in text
    for d in case.phones:
        assert d["label"] in text
    assert "00008120-001A2B3C4D5E801E" in text  # the UDID, written out in full
    assert "How this summary was made" in _text(reader.pages[-1])


def test_pdf_keeps_names_in_other_scripts(demo_case, tmp_path):
    pypdf = pytest.importorskip("pypdf")
    from connected_phone_forensics.output.pdf import write_pdf

    case, _ = demo_case
    case = CaseData.load(case.outdir)
    case.devices[0]["label"] = "Телефон Ивана"
    out = write_pdf(case, tmp_path / "s.pdf")
    text = "\n".join(_text(p) for p in pypdf.PdfReader(str(out)).pages)
    assert "Телефон Ивана" in text


def test_outputs_can_be_skipped(tmp_path):
    from connected_phone_forensics.cli import main

    ev = build_demo_evidence(tmp_path / "demo")
    assert main(["scan", str(ev), "-o", str(tmp_path / "out"), "--no-pdf", "--no-xlsx", "-q"]) == 0
    assert (tmp_path / "out" / "report.html").is_file()
    assert not (tmp_path / "out" / "summary.pdf").exists() and not (tmp_path / "out" / "workbook.xlsx").exists()


def test_view_export_writes_the_rows_shown(demo_case, tmp_path):
    import csv

    from openpyxl import load_workbook

    from connected_phone_forensics.output.excel import write_view, write_view_csv

    case, _ = demo_case
    items = content.extract(case.records, case.devices)
    cat = content.BY_KEY["messages"]
    rows = [r for r in items["messages"] if r["status"]]
    xl = write_view(tmp_path / "v.xlsx", cat, rows, title="Messages · recovered", case=case)
    ws = load_workbook(xl).active
    assert len(ws.tables) == 1 and ws.freeze_panes
    with open(write_view_csv(tmp_path / "v.csv", cat, rows), encoding="utf-8-sig") as f:
        lines = list(csv.reader(f))
    assert lines[0][:3] == ["Time (UTC)", "Phone", "Direction"] and len(lines) == 1 + len(rows)
    assert Path(xl).stat().st_size > 0
