"""Excel workbook: a summary of every phone, and one sheet per phone with what was recovered.

- **Summary**: the case, recovered data at a glance, a table of the phones (identity, how and
  when each connected, how much of each kind was recovered, with a link to its sheet), and the
  computers examined.
- **One sheet per phone**: its identifiers with what each one tells, how it connected, then a
  table per kind of data (messages, calls, contacts, photos and files, notifications, apps,
  connection records), each with the source file and record locator of every row.

Every block is an Excel table with its own sort and filter buttons. Values stay values: times
are real date-times (UTC) so they sort and filter, and evidence text that starts with '=' is
stored as text, never as a formula. Characters Excel cannot store are shown as ``\\xNN``.
Recovered (deleted) rows are shaded amber.
"""

from __future__ import annotations

import csv
import re
from collections import Counter
from collections.abc import Iterable, Sequence
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from .. import __version__, content
from ..casedata import CaseData, os_label
from ..content import CATEGORIES, Category, Column
from .structured import csv_safe

NAVY, NAVY_2, BRAND, YELLOW = "0B1630", "172C55", "2456DE", "FFC23D"
AMBER_SOFT, LINE, MUTED, INK, WHITE = "FFF3D1", "C9D3E6", "5B6B88", "0E1A33", "FFFFFF"
MAX_CELL = 32_000
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


def _styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color=LINE)
    return {
        "title_font": Font(name="Calibri", size=16, bold=True, color=WHITE),
        "sub_font": Font(name="Calibri", size=10, color=WHITE),
        "title_fill": PatternFill("solid", fgColor=NAVY),
        "sub_fill": PatternFill("solid", fgColor=NAVY_2),
        "section_font": Font(name="Calibri", size=12, bold=True, color=INK),
        "section_border": Border(bottom=Side(style="medium", color=YELLOW)),
        "note_font": Font(name="Calibri", size=9, italic=True, color=MUTED),
        "head_font": Font(name="Calibri", size=10, bold=True, color=WHITE),
        "head_fill": PatternFill("solid", fgColor=BRAND),
        "head_align": Alignment(horizontal="center", vertical="center", wrap_text=True),
        "cell_align": Alignment(vertical="top", wrap_text=False),
        "long_align": Alignment(vertical="top", wrap_text=True),
        "num_align": Alignment(horizontal="right", vertical="top"),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "recovered_fill": PatternFill("solid", fgColor=AMBER_SOFT),
        "link_font": Font(name="Calibri", size=10, color=BRAND, underline="single"),
        "kpi_font": Font(name="Calibri", size=10, bold=True, color=INK),
    }


def clean_text(value: Any) -> str:
    text = str(value)
    text = _ILLEGAL.sub(lambda m: f"\\x{ord(m.group()):02x}", text)
    if len(text) > MAX_CELL:
        text = text[:MAX_CELL] + f"… [{len(text):,} characters; full value in timeline.jsonl]"
    return text


def _excel_time(value: Any) -> datetime | None:
    if not value:
        return None
    try:
        dt = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    except ValueError:
        return None
    if dt.tzinfo is not None:
        dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
    return dt


class SheetWriter:
    """Writes title rows, sections, notes and tables down one worksheet."""

    def __init__(self, wb, ws, table_prefix: str):
        self.wb, self.ws = wb, ws
        self.s = _styles()
        self.row = 1
        self.prefix = re.sub(r"\W", "", table_prefix) or "T"
        self.widths: dict[int, float] = {}
        self.max_cols = 1

    # -- blocks -----------------------------------------------------------------------

    def title(self, text: str, subtitle: str = "", span: int = 10) -> None:
        from openpyxl.styles import Alignment

        ws, s = self.ws, self.s
        span = max(span, self.max_cols)
        for r, (value, font, fill, height) in enumerate(((text, s["title_font"], s["title_fill"], 30),
                                                          (subtitle, s["sub_font"], s["sub_fill"], 20))):
            if r == 1 and not subtitle:
                continue
            row = self.row
            for c in range(1, span + 1):
                ws.cell(row=row, column=c).fill = fill
            cell = ws.cell(row=row, column=1, value=clean_text(value))
            cell.font = font
            cell.alignment = Alignment(vertical="center", indent=1)
            ws.row_dimensions[row].height = height
            self.row += 1
        self.row += 1

    def section(self, text: str, span: int = 6) -> None:
        cell = self.ws.cell(row=self.row, column=1, value=clean_text(text))
        cell.font = self.s["section_font"]
        for c in range(1, max(span, 2) + 1):
            self.ws.cell(row=self.row, column=c).border = self.s["section_border"]
        self.ws.row_dimensions[self.row].height = 20
        self.row += 1

    def note(self, text: str) -> None:
        cell = self.ws.cell(row=self.row, column=1, value=clean_text(text))
        cell.font = self.s["note_font"]
        self.row += 1

    def gap(self, n: int = 1) -> None:
        self.row += n

    def table(self, name: str, columns: Sequence[Column], rows: Iterable[dict], *,
              links: dict[str, Any] | None = None, flag: str | None = "status") -> int:
        """An Excel table of ``rows`` (dicts keyed by column key). Returns the number of rows."""
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.table import Table, TableStyleInfo

        ws, s = self.ws, self.s
        top = self.row
        for c, col in enumerate(columns, 1):
            cell = ws.cell(row=top, column=c, value=col.title)
            cell.font, cell.fill, cell.alignment, cell.border = s["head_font"], s["head_fill"], s["head_align"], s["border"]
            self.widths[c] = max(self.widths.get(c, 0), col.width)
        ws.row_dimensions[top].height = 30
        n = 0
        for item in rows:
            n += 1
            r = top + n
            recovered = bool(flag and item.get(flag))
            for c, col in enumerate(columns, 1):
                cell = ws.cell(row=r, column=c)
                self._put(cell, item.get(col.key), col)
                cell.border = s["border"]
                if recovered:
                    cell.fill = s["recovered_fill"]
                link = (links or {}).get(col.key)
                if link and item.get(col.key):
                    target = link(item) if callable(link) else link
                    if target:
                        cell.hyperlink = target
                        cell.font = s["link_font"]
        if n == 0:
            self.row = top + 1
            ws.cell(row=self.row, column=1, value="None").font = s["note_font"]
            self.row += 2
            return 0
        self.max_cols = max(self.max_cols, len(columns))
        ref = f"A{top}:{get_column_letter(len(columns))}{top + n}"
        table = Table(displayName=self._table_name(name), ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        self.row = top + n + 2
        return n

    def _table_name(self, name: str) -> str:
        base = f"{self.prefix}_{re.sub(r'[^A-Za-z0-9_]', '_', name)}"[:240]
        used = self.wb._cpf_table_names  # noqa: SLF001 - one registry per workbook
        candidate, i = base, 2
        while candidate.lower() in used:
            candidate, i = f"{base}_{i}", i + 1
        used.add(candidate.lower())
        return candidate

    def _put(self, cell, value: Any, col: Column) -> None:
        s = self.s
        if col.kind == "time":
            dt = _excel_time(value)
            cell.value = dt
            if dt is not None:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            cell.alignment = s["cell_align"]
            return
        if col.kind == "number" and value not in (None, ""):
            try:
                cell.value = int(value) if float(value).is_integer() else float(value)
                cell.alignment = s["num_align"]
                return
            except (TypeError, ValueError):
                pass
        if value is None:
            cell.value = None
        elif isinstance(value, bool):
            cell.value = "Yes" if value else "No"
        elif isinstance(value, (int, float)):
            cell.value = value
        else:
            if isinstance(value, (list, tuple, set)):
                value = ", ".join(str(v) for v in value)
            text = clean_text(value)
            cell.value = text
            cell.data_type = "s"  # never a formula, whatever the evidence says
        cell.alignment = s["long_align"] if col.kind == "long" else s["cell_align"]

    def finish(self, freeze: str | None = None, landscape: bool = True) -> None:
        from openpyxl.utils import get_column_letter

        ws = self.ws
        for c, w in self.widths.items():
            ws.column_dimensions[get_column_letter(c)].width = w + 2
        if freeze:
            ws.freeze_panes = freeze
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.oddFooter.center.text = "Connected Phone Forensics · &A · page &P of &N"


def _new_workbook(title: str, subject: str):
    from openpyxl import Workbook

    wb = Workbook()
    wb._cpf_table_names = set()  # noqa: SLF001
    wb.properties.title = title
    wb.properties.subject = subject
    wb.properties.creator = f"Connected Phone Forensics {__version__}"
    return wb


def sheet_title(text: str, used: set[str]) -> str:
    name = _SHEET_BAD.sub("-", text).strip().strip("'") or "Phone"
    name = name[:31]
    candidate, i = name, 2
    while candidate.lower() in used or candidate.lower() == "summary":
        suffix = f" ({i})"
        candidate, i = name[:31 - len(suffix)] + suffix, i + 1
    used.add(candidate.lower())
    return candidate


def _case_line(case: CaseData) -> str:
    parts = []
    if case.case.get("case"):
        parts.append(f"Case {case.case['case']}")
    if case.case.get("examiner"):
        parts.append(f"Examiner {case.case['examiner']}")
    if case.started:
        parts.append(f"Scanned {content.format_time(case.started)} UTC")
    parts.append(f"Connected Phone Forensics {case.tool_version or __version__}")
    return " · ".join(parts)


SUMMARY_PHONE_COLUMNS = (
    Column("label", "Phone", 26), Column("id", "Record", 10), Column("os", "OS", 12), Column("model", "Model", 26),
    Column("status", "Identification", 18), Column("primary_label", "Main identifier", 18),
    Column("primary_value", "Value", 30), Column("primary_summary", "What it tells", 40, "long"),
    Column("numbers", "Phone numbers", 18), Column("computers", "Computers", 22),
    Column("first_seen", "First seen (UTC)", 19, "time"), Column("last_seen", "Last seen (UTC)", 19, "time"),
    Column("connections", "Connected by", 24),
    *(Column(c.key, c.title, 11, "number") for c in CATEGORIES),
    Column("recovered", "Recovered (deleted)", 12, "number"), Column("sheet", "Details", 14),
)
_IDENTIFIER_COLUMNS = (
    Column("label", "Identifier", 20), Column("value", "Value", 34), Column("meaning", "What it tells", 46, "long"),
    Column("details", "Decoded", 46, "long"), Column("strength", "Strength", 10),
    Column("sources", "Recorded by", 30),
)
_CHANNEL_COLUMNS = (
    Column("channel", "Connection", 18), Column("first_seen", "First seen (UTC)", 19, "time"),
    Column("last_seen", "Last seen (UTC)", 19, "time"), Column("observations", "Records", 10, "number"),
    Column("artifacts", "Parsers", 40),
)
_HOTSPOT_COLUMNS = (
    Column("ssid", "Network name (SSID)", 28), Column("likely", "Likely phone (from the name)", 26),
    Column("access_points", "Access points", 12, "number"), Column("randomised", "Randomised addresses", 13, "number"),
    Column("first", "First connected (UTC)", 19, "time"), Column("last", "Last connected (UTC)", 19, "time"),
    Column("computers", "Computers", 22), Column("sheet", "Details", 14),
)
_SESSION_COLUMNS = (
    Column("bssid", "Access point (BSSID)", 20), Column("maker", "Maker", 30),
    Column("first", "First connected (UTC)", 19, "time"), Column("last", "Last connected (UTC)", 19, "time"),
    Column("profiles", "Windows profile", 26), Column("names", "Network names", 30), Column("dhcp", "Addresses", 30),
    Column("computers", "Computer", 16), Column("device", "Record", 10), Column("sources", "Source", 48, "path"),
)
_SYSTEM_COLUMNS = (Column("label", "Computer", 22), Column("os", "OS", 10), Column("version", "Version", 40),
                   Column("users", "User profiles", 30))
_KPI_COLUMNS = (Column("title", "Recovered", 18), Column("items", "Items", 10, "number"),
                Column("phones", "Phones", 10, "number"), Column("recovered", "Recovered (deleted)", 14, "number"),
                Column("description", "What it covers", 50))


def write_workbook(case: CaseData, path: Path, items: dict[str, list[dict]] | None = None) -> Path:
    """The whole case: a Summary sheet, then one sheet per phone."""
    path = Path(path)
    items = items if items is not None else content.extract(case.records, case.devices)
    per_device = content.counts(items)
    title = f"Connected Phone Forensics — {case.title}"
    wb = _new_workbook(title, case.case.get("case") or "")
    summary = wb.active
    summary.title = "Summary"
    summary.sheet_properties.tabColor = NAVY

    used: set[str] = set()
    phones = case.phones
    labels = Counter(d.get("label") or d["id"] for d in phones)

    def sheet_for(d: dict) -> str:
        lbl = d.get("label") or d["id"]
        short = lbl.split(" · ")[0]  # "iPhone 16 family · UDID 0000…" -> "iPhone 16 family"
        if labels[lbl] > 1 or short != lbl or len(lbl) > 31:  # the record ID tells such sheets apart
            lbl = f"{short[:31 - len(d['id']) - 3].rstrip()} ({d['id']})"
        return sheet_title(lbl, used)

    sheet_names = {d["id"]: sheet_for(d) for d in phones}
    unlinked = per_device.get(None)
    if unlinked:
        sheet_names[None] = sheet_title(content.UNLINKED, used)

    w = SheetWriter(wb, summary, "Summary")
    w.title("Connected Phone Forensics — case summary", _case_line(case), span=len(SUMMARY_PHONE_COLUMNS))
    w.note(f"Evidence: {case.input_path}")
    if case.acquisition:
        a = case.acquisition
        w.note(f"Live acquisition of {a.get('computer')} ({a.get('windows') or 'Windows'}) by {a.get('user')}: "
               f"{a.get('files_acquired')} files and {a.get('hives_saved')} registry hives copied with their "
               f"SHA-256, {a.get('errors')} error(s); started {content.format_time(a.get('started'))} UTC.")
    w.note("All times are UTC. Unconfirmed phones are identified only by weak identifiers (names, numbers).")
    w.gap()

    w.section("Recovered data at a glance", span=len(_KPI_COLUMNS))
    kpis = []
    for c in CATEGORIES:
        rows = items[c.key]
        kpis.append({"title": c.title, "items": len(rows), "phones": len({r["device"] for r in rows if r["device"]}),
                     "recovered": sum(1 for r in rows if r.get("status")), "description": c.description})
    kpis.append({"title": "Total", "items": sum(k["items"] for k in kpis),
                 "phones": len({r["device"] for rows in items.values() for r in rows if r["device"]}),
                 "recovered": sum(k["recovered"] for k in kpis), "description": ""})
    w.table("glance", _KPI_COLUMNS, kpis, flag=None)

    w.section(f"Phones ({len(phones)})", span=len(SUMMARY_PHONE_COLUMNS))
    phone_rows = []
    for d in phones:
        facts = content.phone_facts(d)
        c = per_device.get(d["id"], Counter())
        phone_rows.append({**facts, "numbers": ", ".join(facts["numbers"]), "computers": ", ".join(facts["computers"]),
                           "first_seen": d.get("first_seen"), "last_seen": d.get("last_seen"),
                           "connections": ", ".join(facts["connections"]),
                           **{cat.key: c.get(cat.key, 0) for cat in CATEGORIES},
                           "recovered": sum(1 for rows in items.values() for r in rows
                                            if r["device"] == d["id"] and r.get("status")),
                           "sheet": "Open sheet", "_sheet": sheet_names[d["id"]]})
    w.table("phones", SUMMARY_PHONE_COLUMNS, phone_rows, flag=None,
            links={"sheet": lambda r: f"#'{r['_sheet']}'!A1"})
    if unlinked:
        w.note(f"{sum(unlinked.values())} item(s) could not be linked to a phone (for example Phone Link "
               f"notifications without a device): see the sheet '{sheet_names[None]}'.")
        w.gap()

    hotspots = case.hotspots
    if hotspots:
        hotspot_sheet = sheet_title("Hotspots", used)
        w.section(f"Phone hotspots joined ({len(hotspots)} networks)", span=len(_HOTSPOT_COLUMNS))
        w.note("Phones known only as Wi-Fi hotspots, by network name. A phone that randomises its hotspot address "
               "appears as several access points with the same name; each is listed on the Hotspots sheet.")
        w.table("hotspots", _HOTSPOT_COLUMNS, [
            {"ssid": n["ssid"], "likely": n["likely_phone"] or {"ios": "iPhone", "android": "Android phone"}.get(
                n["os"], ""), "access_points": len(n["sessions"]), "randomised": n["randomised"],
             "first": n["first"], "last": n["last"], "computers": ", ".join(n["computers"]),
             "sheet": "Open sheet"} for n in hotspots], flag=None,
            links={"sheet": f"#'{hotspot_sheet}'!A1"})

    w.section(f"Computers examined ({len(case.systems)})", span=len(_SYSTEM_COLUMNS))
    w.table("computers", _SYSTEM_COLUMNS, [_system_row(s) for s in case.systems], flag=None)
    w.finish(freeze="A4")

    for d in phones:
        _phone_sheet(wb, sheet_names[d["id"]], d, items, outdir=case.outdir, target=path)
    if hotspots:
        _hotspot_sheet(wb, hotspot_sheet, hotspots)
    if unlinked:
        _phone_sheet(wb, sheet_names[None], None, items, outdir=case.outdir, target=path)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _system_row(s: dict) -> dict:
    det = s.get("details") or {}
    version = det.get("product") or det.get("pretty_name") or (
        f"macOS {det['product_version']}" if det.get("product_version") else "")
    extra = " ".join(str(v) for k, v in det.items()
                     if k in ("display_version", "build", "edition") and v)
    return {"label": s.get("label"), "os": os_label(s.get("os")), "version": f"{version} {extra}".strip(),
            "users": ", ".join(s.get("users") or [])}


def _file_link(outdir: Path | None, target: Path):
    def link(item: dict) -> str | None:
        rel = item.get("file")
        if not rel or outdir is None:
            return None
        exported = Path(outdir) / "exports" / rel
        try:
            return exported.relative_to(target.parent).as_posix()
        except ValueError:
            return exported.as_uri() if exported.is_absolute() else None
    return link


def _phone_sheet(wb, name: str, device: dict | None, items: dict[str, list[dict]], *, outdir, target: Path) -> None:
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = BRAND if device else "8A94A6"
    w = SheetWriter(wb, ws, name)
    dev_id = device["id"] if device else None
    if device:
        facts = content.phone_facts(device)
        line = " · ".join(x for x in (facts["os"], facts["model"], facts["status"],
                                      f"computers: {', '.join(facts['computers'])}" if facts["computers"] else "",
                                      f"first seen {facts['first_seen']} UTC" if facts["first_seen"] else "",
                                      f"last seen {facts['last_seen']} UTC" if facts["last_seen"] else "") if x)
        w.title(facts["label"], f"{line} · record {device['id']}", span=10)
        w.section("Identifiers", span=len(_IDENTIFIER_COLUMNS))
        w.table("identifiers", _IDENTIFIER_COLUMNS, [
            {**i, "details": "; ".join(f"{k}: {v}" for k, v in i["fields"]) + (f" — {i['note']}" if i["note"] else ""),
             "sources": ", ".join(i["sources"])} for i in facts["identifiers"]], flag=None)
        w.section("How it connected", span=len(_CHANNEL_COLUMNS))
        w.table("channels", _CHANNEL_COLUMNS, [
            {**c, "channel": content.CHANNEL_LABEL.get(c["channel"], c["channel"]),
             "artifacts": ", ".join(c["artifacts"])} for c in device.get("channels") or []], flag=None)
    else:
        w.title(content.UNLINKED, "Recovered items that no identifier ties to a particular phone", span=10)
    any_content = False
    for cat in CATEGORIES:
        rows = [r for r in items[cat.key] if r["device"] == dev_id]
        if not rows:
            continue
        any_content = any_content or cat.key != "connections"
        cols = [c for c in cat.columns if c.key != "phone"]
        w.section(f"{cat.title} ({len(rows)})", span=len(cols))
        w.table(cat.key, cols, rows, links={"file": _file_link(outdir, target)})
    if device and not any_content:
        w.note("No messages, calls, contacts, photos, notifications or apps were recovered for this phone: the "
               "computers only recorded how it connected.")
    w.finish(freeze="A4")


def _hotspot_sheet(wb, name: str, hotspots: list[dict]) -> None:
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = "0E8FB5"
    w = SheetWriter(wb, ws, "Hotspots")
    total = sum(len(n["sessions"]) for n in hotspots)
    w.title("Phone hotspots joined", f"{len(hotspots)} network names · {total} access points", span=10)
    w.note("Android 11 and later, and recent iPhones, give the hotspot a new random address (BSSID) each time it is "
           "switched on: one phone's hotspot then appears as several access points with the same name. A shared "
           "name does not prove one phone. Randomised addresses name no maker and cannot be located.")
    w.gap()
    for n, net in enumerate(hotspots, 1):
        likely = net["likely_phone"] or {"ios": "iPhone", "android": "Android phone"}.get(net["os"], "")
        w.section(f"“{net['ssid']}” · {len(net['sessions'])} access point(s)" + (f" · likely {likely}" if likely
                                                                                 else ""),
                  span=len(_SESSION_COLUMNS))
        w.table(f"network_{n}", _SESSION_COLUMNS,
                [{**s, "sources": ", ".join(s["sources"])} for s in net["sessions"]], flag=None)
    w.finish(freeze="A4")


def write_view(path: Path, category: Category, rows: list[dict], *, title: str, subtitle: str = "",
               case: CaseData | None = None) -> Path:
    """The rows a dashboard view shows, as one Excel table."""
    path = Path(path)
    wb = _new_workbook(title, case.case.get("case") if case else "")
    ws = wb.active
    ws.title = sheet_title(category.title, set())
    w = SheetWriter(wb, ws, category.key)
    sub = " · ".join(x for x in (subtitle, _case_line(case) if case else "") if x)
    w.title(title, sub, span=len(category.columns))
    header_row = w.row
    w.table(category.key, category.columns, rows,
            links={"file": _file_link(case.outdir if case else None, path)})
    w.finish(freeze=f"A{header_row + 1}")
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def write_view_csv(path: Path, category: Category, rows: list[dict]) -> Path:
    path = Path(path)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        out = csv.writer(f)
        out.writerow([c.title for c in category.columns])
        for r in rows:
            out.writerow([csv_safe(content.cell(r, c)) for c in category.columns])
    return path
